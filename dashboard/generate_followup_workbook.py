from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from analysis import FOLLOWUP_OFFSETS, followup_schedule, load_and_analyse


APP_DIR = Path(__file__).resolve().parent
PROJECT_DIR = APP_DIR.parent
DEFAULT_DATA = PROJECT_DIR / "DATA" / "Excel_DATA-1.0-2026_08_26_09_49_19.xlsx"
DEFAULT_OUTPUT = PROJECT_DIR / "outputs" / "followup_schedule" / "XEAS随访表_20260826.xlsx"


def create_followup_workbook(data_file: Path = DEFAULT_DATA, output_file: Path = DEFAULT_OUTPUT) -> Path:
    bundle = load_and_analyse(data_file, data_file.name)
    schedule = followup_schedule(bundle)
    output_file.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "随访表"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "G4"
    ws.auto_filter.ref = f"A3:Z{len(schedule) + 3}"

    blue = "8DB4E2"
    dark_blue = "17365D"
    pale_blue = "DCE6F1"
    red = "C00000"
    green = "D9EAD3"
    yellow = "FFF2CC"
    gray = "E7E6E6"
    thin = Side(style="thin", color="A6A6A6")

    ws.merge_cells("A1:Z1")
    ws["A1"] = "XEAS项目随访表"
    ws["A1"].fill = PatternFill("solid", fgColor=blue)
    ws["A1"].font = Font(name="微软雅黑", size=18, bold=True, color=dark_blue)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34

    fixed = [("A2:A3", "受试者编号"), ("B2:B3", "年龄"), ("C2:C3", "年龄组"), ("D2:D3", "受试者状态"),
             ("E2:E3", "知情日期\nV1（筛选，不超过2周）"), ("F2:F3", "V1发药日期\nV1（基线 D-1）")]
    for cell_range, label in fixed:
        ws.merge_cells(cell_range)
        cell = ws[cell_range.split(":")[0]]
        cell.value = label

    visit_starts = {"V2": 7, "V3": 11, "V4": 15, "V5": 19, "V6": 23}
    for visit, start_col in visit_starts.items():
        end_col = start_col + 3
        ws.merge_cells(start_row=2, start_column=start_col, end_row=2, end_column=end_col)
        ws.cell(2, start_col, f"{visit}（D{FOLLOWUP_OFFSETS[visit]}±2天）")
        for offset, label in enumerate(["理论日期", "下限", "上限", "实际访视日期"]):
            ws.cell(3, start_col + offset, label)

    for row in ws.iter_rows(min_row=2, max_row=3, min_col=1, max_col=26):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.font = Font(name="微软雅黑", size=10, bold=True, color=dark_blue)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in [10, 14, 18, 22, 26]:
        ws.cell(3, col).font = Font(name="微软雅黑", size=10, bold=True, color=red)
    ws.row_dimensions[2].height = 25
    ws.row_dimensions[3].height = 32

    for idx, record in schedule.iterrows():
        row = idx + 4
        ws.cell(row, 1, record["受试者编号"])
        ws.cell(row, 2, None if record["年龄"] != record["年龄"] else float(record["年龄"]))
        ws.cell(row, 3, record["年龄组"])
        ws.cell(row, 4, record["受试者状态"])
        ws.cell(row, 5, record["知情日期"].to_pydatetime() if record["知情日期"] == record["知情日期"] else None)
        ws.cell(row, 6, record["V1发药日期"].to_pydatetime() if record["V1发药日期"] == record["V1发药日期"] else None)
        for visit, start_col in visit_starts.items():
            days = FOLLOWUP_OFFSETS[visit]
            actual = record[f"{visit}实际日期"]
            ws.cell(row, start_col, f'=IF(OR($F{row}="",$D{row}="筛选失败"),"",$F{row}+{days})')
            ws.cell(row, start_col + 1, f'=IF({get_column_letter(start_col)}{row}="","",{get_column_letter(start_col)}{row}-2)')
            ws.cell(row, start_col + 2, f'=IF({get_column_letter(start_col)}{row}="","",{get_column_letter(start_col)}{row}+2)')
            ws.cell(row, start_col + 3, actual.to_pydatetime() if actual == actual else None)
        for cell in ws[row]:
            cell.font = Font(name="微软雅黑", size=9)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for col in range(5, 27):
            ws.cell(row, col).number_format = "yyyy-mm-dd"
        ws.cell(row, 2).number_format = "0.000"

    last_row = len(schedule) + 3
    ws.conditional_formatting.add(f"A4:Z{last_row}", FormulaRule(formula=["$D4=\"筛选失败\""], fill=PatternFill("solid", fgColor=yellow)))
    ws.conditional_formatting.add(f"A4:Z{last_row}", FormulaRule(formula=["$D4=\"退出\""], fill=PatternFill("solid", fgColor=green)))
    ws.conditional_formatting.add(f"A4:Z{last_row}", FormulaRule(formula=["$D4=\"数据待补充\""], fill=PatternFill("solid", fgColor=gray)))

    widths = {"A": 13, "B": 9, "C": 17, "D": 12, "E": 17, "F": 15}
    for col in range(7, 27):
        widths[get_column_letter(col)] = 13
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:3"

    notes = wb.create_sheet("说明")
    notes.sheet_view.showGridLines = False
    note_rows = [
        ["项目", "说明"],
        ["数据来源", data_file.name + " 的SV及DM工作表"],
        ["隐私处理", "未输出姓名、门诊号、电话等直接身份信息。"],
        ["知情日期", "SV工作表中SCR（V1筛选，不超过2周）的实际访视日期。"],
        ["V1发药日期", "SV工作表中V1（基线 D-1）的实际访视日期。"],
        ["理论访视", "以V1发药日期为起点：V2=D14、V3=D28、V4=D42、V5=D56、V6=D84。"],
        ["允许窗", "理论访视日期前后各2天。"],
        ["颜色", "黄色=筛选失败；绿色=退出；灰色=数据待补充。"],
        ["筛选失败", "仅保留筛选期知情日期，不生成发药及后续访视计划。"],
    ]
    for row in note_rows:
        notes.append(row)
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 80
    notes.freeze_panes = "A2"
    for cell in notes[1]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(name="微软雅黑", bold=True, color=dark_blue)
    for row in notes.iter_rows():
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, bold=cell.row == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color="D9E2F3"))
    notes.row_dimensions[1].height = 24

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(output_file)
    return output_file


if __name__ == "__main__":
    print(create_followup_workbook())
